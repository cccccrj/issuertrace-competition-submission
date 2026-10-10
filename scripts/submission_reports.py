"""Render formal submission Markdown to DOCX with portable, clickable references.

Tool-only dependency: python-docx. PDF conversion uses LibreOffice separately.
Relative references resolve to the chosen GitHub revision, not a local
machine path. Set --ref to the frozen tag when producing a final release.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def plain(text: str) -> str:
    return text.replace("**", "").replace("`", "")


def reference_url(target: str, source: Path, ref: str) -> str:
    target = target.strip("<>")
    parsed = urlsplit(target)
    if parsed.scheme:
        if parsed.scheme not in {"http", "https", "mailto"}:
            raise ValueError("Unsupported reference scheme")
        return target
    path = (source.parent / unquote(parsed.path)).resolve() if parsed.path else source.resolve()
    relative = path.relative_to(ROOT).as_posix()
    repo = ROOT.name
    if repo not in {"issuertrace-competition-submission", "courseloop-competition-submission"}:
        raise ValueError("Render from the submission source repository")
    kind = "tree" if path.is_dir() else "blob"
    url = f"https://github.com/cccccrj/{repo}/{kind}/{quote(ref, safe='/')}/{quote(relative, safe='/')}"
    return url + ("#" + quote(unquote(parsed.fragment), safe="-_") if parsed.fragment else "")


def add_text(paragraph, text: str, source: Path, ref: str) -> None:
    position = 0
    for match in LINK.finditer(text):
        paragraph.add_run(plain(text[position:match.start()]))
        url = reference_url(match.group(2), source, ref)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.extend([color, underline])
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = plain(match.group(1))
        run.append(value)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        position = match.end()
    paragraph.add_run(plain(text[position:]))


def render(source: Path, ref: str = "v1.0") -> Path:
    doc = Document()
    doc.core_properties.author = "Submission project contributors"
    doc.core_properties.last_modified_by = "Submission report generator"
    section = doc.sections[0]
    project = "IssuerTrace" if ROOT.name.startswith("issuertrace") else "CourseLoop"
    section.header.paragraphs[0].text = project + " · 比赛交付材料"
    section.left_margin = section.right_margin = Inches(.8)
    style = doc.styles["Normal"]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Droid Sans Fallback")
    style.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        s = doc.styles[name]
        s.font.color.rgb = RGBColor.from_string("253653")
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Droid Sans Fallback")
    footer = section.footer.paragraphs[0]
    footer.text = "正式提交版本 V1.0    "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    rows: list[list[str]] = []
    code = False

    def flush() -> None:
        if not rows:
            return
        table = doc.add_table(rows=1, cols=max(len(r) for r in rows))
        table.style = "Light Shading Accent 1"
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for i, row in enumerate(rows):
            current = table.rows[0] if i == 0 else table.add_row()
            current._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            cells = current.cells
            for j, value in enumerate(row):
                add_text(cells[j].paragraphs[0], value, source, ref)
        rows.clear()

    for line in source.read_text().splitlines():
        if line.startswith("|") and not code:
            row = [x.strip() for x in re.split(r"(?<!\\)\|", line.strip("|"))]
            if not all(re.fullmatch(r"[:\- ]+", x or "-") for x in row):
                rows.append(row)
            continue
        flush()
        if line.startswith("```"):
            code = not code
            continue
        if not line.strip():
            continue
        if code:
            paragraph = doc.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(8)
                run.font.name = "DejaVu Sans Mono"
            continue
        match = IMAGE.fullmatch(line)
        if match:
            path = (source.parent / unquote(match.group(2))).resolve()
            path.relative_to(ROOT)
            if not path.is_file():
                raise FileNotFoundError(path)
            doc.add_picture(str(path), width=Inches(6.4))
            doc.add_paragraph(match.group(1), "Caption")
            continue
        heading = re.match(r"^(#{1,4}) (.*)$", line)
        if heading:
            paragraph = doc.add_heading("", min(len(heading.group(1)) - 1, 3))
            add_text(paragraph, heading.group(2), source, ref)
        elif line.startswith("- "):
            add_text(doc.add_paragraph(style="List Bullet"), line[2:], source, ref)
        elif line.startswith("> "):
            add_text(doc.add_paragraph(style="Quote"), line[2:], source, ref)
        else:
            add_text(doc.add_paragraph(), line, source, ref)
    flush()
    dest = source.with_suffix(".docx")
    doc.save(dest)
    return dest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ref", default="v1.0",
                        help="GitHub revision for portable report links; use the final tag at freeze")
    args = parser.parse_args()
    for source in sorted((ROOT / "docs/submission").glob("[0-9][0-9]-*.md")):
        print(render(source, args.ref).name)


if __name__ == "__main__":
    main()
